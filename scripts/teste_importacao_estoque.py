"""Prova da importação de peças por planilha (Padrao_Excel_Template_Importacao).

Roda em banco DESCARTÁVEL, nunca no instance/oficina.db.
    python scripts/teste_importacao_estoque.py
Sai com código 1 se algum caso falhar.
"""
import io
import os
import re
import sys
import tempfile

os.environ["FLASK_DATABASE"] = os.path.join(tempfile.mkdtemp(prefix="oficina_teste_"), "t.db")
os.environ["ADMIN_EMAIL"] = "admin@teste.local"
os.environ["ADMIN_SENHA_INICIAL"] = "SenhaTesteAdmin2026"
os.environ.pop("ADMIN_RESET_SENHA", None)

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import bcrypt  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from app import create_app  # noqa: E402

app = create_app()
falhas = []


def chk(rotulo, cond, detalhe=""):
    print(("OK   " if cond else "FAIL ") + rotulo + ("" if cond else f"  <- {detalhe}"))
    if not cond:
        falhas.append(rotulo)


def sql(q, p=()):
    with app.app_context():
        from core.db import get_db
        db = get_db()
        r = db.execute(q, p).fetchall()
        db.commit()
        return r


def planilha(linhas, cabecalho=True):
    w = Workbook()
    s = w.active
    s.title = "Peças"
    if cabecalho:
        s.append(["Codigo", "Descricao", "Unidade", "Preco Custo", "Preco Venda",
                  "Estoque Atual", "Estoque Minimo", "Localizacao"])
    for ln in linhas:
        s.append(ln)
    buf = io.BytesIO()
    w.save(buf)
    buf.seek(0)
    return buf


def peca(codigo):
    r = sql("SELECT descricao, unidade, preco_custo, preco_venda, estoque_atual FROM peca WHERE codigo = ? AND D_E_L_E_T = 0", (codigo,))
    if not r:
        return None
    d, u, pc, pv, ea = r[0]
    return (d, u, float(pc), float(pv), float(ea))   # SQLite NUMERIC devolve número


sql("UPDATE usuario SET precisa_trocar_senha = 0")
h = bcrypt.hashpw(b"SenhaTecnico2026", bcrypt.gensalt()).decode()
sql("INSERT INTO usuario (nome, email, senha_hash, id_perfil, precisa_trocar_senha) "
    "VALUES ('Tec', 'tec@teste.local', ?, 2, 0)", (h,))

adm = app.test_client()
adm.post("/auth/login", data={"email": "admin@teste.local", "senha": "SenhaTesteAdmin2026"})
TK = re.search(rb'name="csrf-token" content="([^"]+)"', adm.get("/estoque/").data).group(1).decode()
tec = app.test_client()
tec.post("/auth/login", data={"email": "tec@teste.local", "senha": "SenhaTecnico2026"})


def subir(c, buf, nome="pecas.xlsx", tk=None):
    return c.post("/estoque/importar", data={"csrf_token": tk or TK, "planilha": (buf, nome)},
                  content_type="multipart/form-data")


def confirmar(c, dados, tk=None):
    return c.post("/estoque/importar/confirmar", data={"csrf_token": tk or TK, "dados": dados})


def dados_da_previa(resp):
    m = re.search(rb'name="dados" value="([^"]*)"', resp.data)
    import html
    return html.unescape(m.group(1).decode()) if m else None


# ── 1. Modelo ────────────────────────────────────────────────────────────────
r = adm.get("/estoque/modelo.xlsx")
chk("modelo baixa como .xlsx", r.status_code == 200 and "spreadsheetml" in r.mimetype)
chk("nome do arquivo é descritivo", "modelo_estoque_pecas.xlsx" in r.headers.get("Content-Disposition", ""))
wb = load_workbook(io.BytesIO(r.data))
ws = wb["Peças"]
chk("modelo: cabeçalho na linha 1, exemplo EX-001, freeze A2 e aba de ajuda",
    ws.cell(1, 1).value == "Codigo" and ws.cell(2, 1).value == "EX-001"
    and ws.freeze_panes == "A2" and "Como preencher" in wb.sheetnames)
chk("técnico não baixa o modelo -> 403", tec.get("/estoque/modelo.xlsx").status_code == 403)
chk("técnico não abre a importação -> 403", tec.get("/estoque/importar").status_code == 403)

# ── 2. Prévia não grava nada ─────────────────────────────────────────────────
buf = planilha([
    ["EX-001", "Exemplo", "UN", 1, 2, 3, 4, ""],          # ignorado
    ["RES-220", "Resistência 220V", "UN", 40, 90, 5, 3, "B1"],   # novo
    ["TER-01", "Termostato", "PC", 18.5, 35, 8, 2, "A2"],        # novo
    ["RES-220", "Duplicata", "UN", 1, 1, 1, 1, ""],             # repetido
    ["Y", "curto demais", "UN", 1, 2, 3, 4, ""],               # erro codigo
    ["BOMBA", "Bomba", "LITRO", 1, 2, 3, 4, ""],               # erro unidade
])
r = subir(adm, buf)
html_previa = r.data.decode()
chk("prévia abre com o resumo", r.status_code == 200 and "linha(s) lida(s)" in html_previa)
chk("prévia conta 2 novas, 0 atualiza, 3 erros (exemplo e branca fora)",
    "<b>2</b> nova" in html_previa and "<b>0</b> atualiza" in html_previa and "<b>3</b> com erro" in html_previa, )
chk("prévia NÃO gravou nada no banco", len(sql("SELECT 1 FROM peca")) == 0)
chk("erros aparecem com a linha e o motivo", "Código repetido" in html_previa and "Unidade inválida" in html_previa)

# ── 3. Confirmar grava (upsert por código) ──────────────────────────────────
dados = dados_da_previa(r)
chk("prévia carrega os dados crus pra confirmar", dados and "RES-220" in dados)
r2 = confirmar(adm, dados)
chk("confirmar redireciona pro estoque", r2.status_code == 302 and "/estoque" in r2.headers["Location"])
chk("2 peças criadas (exemplo e erros fora)", len(sql("SELECT 1 FROM peca WHERE D_E_L_E_T = 0")) == 2)
chk("RES-220 gravada com os valores certos (número não vira lixo)",
    peca("RES-220") == ("Resistência 220V", "UN", 40.0, 90.0, 5.0), peca("RES-220"))
chk("TER-01 gravada com decimal certo (18.5)", peca("TER-01")[2] == 18.5, peca("TER-01"))

# ── 4. Idempotente: subir de novo ATUALIZA, não duplica ─────────────────────
buf2 = planilha([["RES-220", "Resistência 220V nova", "UN", 42, 95, 7, 3, "B2"]])
r = subir(adm, buf2)
chk("prévia marca RES-220 como 'atualiza' (já existe)", "<b>1</b> atualiza" in r.data.decode())
confirmar(adm, dados_da_previa(r))
chk("continua 2 peças (não duplicou)", len(sql("SELECT 1 FROM peca WHERE D_E_L_E_T = 0")) == 2)
chk("RES-220 foi atualizada (descrição e preço novos)",
    peca("RES-220") == ("Resistência 220V nova", "UN", 42.0, 95.0, 7.0), peca("RES-220"))

# subir EXATAMENTE o mesmo arquivo 2x não muda mais nada
confirmar(adm, dados_da_previa(subir(adm, planilha([["RES-220", "Resistência 220V nova", "UN", 42, 95, 7, 3, "B2"]]))))
chk("reimportar o mesmo arquivo é idempotente (segue 2 peças)", len(sql("SELECT 1 FROM peca WHERE D_E_L_E_T = 0")) == 2)

# ── 5. Formatos e vírgula decimal ───────────────────────────────────────────
buf3 = planilha([
    ["PBR-1", "Preço pt-BR texto", "UN", "1.234,56", "2.000,00", "1,5", "0", ""],  # texto pt-BR
    ["cx-1", "código minúsculo", "cx", 5, 9, 2, 1, ""],                            # vira maiúsculo
])
confirmar(adm, dados_da_previa(subir(adm, buf3)))
chk("texto pt-BR '1.234,56' vira 1234.56", peca("PBR-1")[2] == 1234.56, peca("PBR-1"))
chk("estoque '1,5' vira 1.5", peca("PBR-1")[4] == 1.5, peca("PBR-1"))
chk("código minúsculo é gravado em maiúsculo", peca("CX-1") is not None and peca("cx-1") is None)

# ── 6. Arquivo inválido e vazio ─────────────────────────────────────────────
r = adm.post("/estoque/importar", data={"csrf_token": TK, "planilha": (io.BytesIO(b"nao sou xlsx"), "fake.txt")},
             content_type="multipart/form-data")
chk("arquivo não-.xlsx é recusado", "Excel .xlsx" in r.data.decode())
r = adm.post("/estoque/importar", data={"csrf_token": TK, "planilha": (io.BytesIO(b"PK\x03\x04 lixo"), "fake.xlsx")},
             content_type="multipart/form-data")
chk("xlsx corrompido é recusado com mensagem", "Excel" in r.data.decode() and r.status_code == 200)
r = subir(adm, planilha([["EX-001", "Só exemplo", "UN", 1, 2, 3, 4, ""]]))
chk("planilha só com o exemplo -> nada pra importar", "não tem nenhuma linha" in r.data.decode()
    or "<b>0</b> nova" in r.data.decode())

# ── 7. Segurança ─────────────────────────────────────────────────────────────
TKT = re.search(rb'name="csrf-token" content="([^"]+)"', tec.get("/home").data).group(1).decode()
r = subir(tec, planilha([["HACK-1", "do técnico", "UN", 1, 2, 3, 4, ""]]), tk=TKT)
chk("técnico não importa -> 403", r.status_code == 403 and peca("HACK-1") is None)
n = len(sql("SELECT 1 FROM peca WHERE D_E_L_E_T = 0"))
r = adm.post("/estoque/importar/confirmar", data={"dados": '[["SEMTK-1","x","UN",1,2,3,4,""]]'})  # sem csrf
chk("confirmar sem CSRF não grava", len(sql("SELECT 1 FROM peca WHERE D_E_L_E_T = 0")) == n and peca("SEMTK-1") is None)
chk("auditoria registra a importação", len(sql("SELECT 1 FROM logs_acesso WHERE acao = 'ESTOQUE_IMPORTADO'")) >= 1)
chk("dev DB nunca é o oficina.db (sanidade)", "oficina.db" not in os.environ["FLASK_DATABASE"])

print(f"\nFALHAS: {len(falhas)}")
sys.exit(1 if falhas else 0)
