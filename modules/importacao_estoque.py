"""Importação de peças por planilha (Padrao_Excel_Template_Importacao).

O usuário PREENCHE um modelo (não vem de sistema externo): o sistema oferece o
download do modelo pronto (cabeçalho colorido, exemplo, instruções, validações
nativas do Excel) e lê de volta. Reaproveita PecaIn pra validar cada linha.

Idempotente: a chave é `peca.codigo` (UNIQUE). Código que já existe ATUALIZA;
novo INSERE. Subir o mesmo arquivo 2x não duplica.
"""
from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.db import query_one
from modules.validacao_peca import UNIDADES, PecaIn

# (título da coluna, chave do campo). A ordem é a ordem das colunas na planilha.
COLUNAS = [
    ("Codigo", "codigo"),
    ("Descricao", "descricao"),
    ("Unidade", "unidade"),
    ("Preco Custo", "preco_custo"),
    ("Preco Venda", "preco_venda"),
    ("Estoque Atual", "estoque_atual"),
    ("Estoque Minimo", "estoque_minimo"),
    ("Localizacao", "localizacao"),
]
VERDE_CADASTRO = "15803D"   # cabeçalho verde: contexto cadastro (Padrao_Excel_Template)
COD_EXEMPLO = "EX-001"      # a linha de exemplo é ignorada na leitura
MAX_LINHAS = 2000           # teto de linhas de dados por arquivo


# ──────────────────────────────────────────────────────────────────────────────
# Geração do modelo
# ──────────────────────────────────────────────────────────────────────────────

def gerar_modelo() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Peças"

    hfill = PatternFill("solid", fgColor=VERDE_CADASTRO)
    hfont = Font(color="FFFFFF", bold=True)
    for i, (titulo, _) in enumerate(COLUNAS, start=1):
        cel = ws.cell(row=1, column=i, value=titulo)
        cel.fill = hfill
        cel.font = hfont
        cel.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(titulo) + 3)

    # Linha 2: exemplo realista (código EX-001 é ignorado na importação)
    for i, v in enumerate([COD_EXEMPLO, "Filtro de óleo", "UN", 12.50, 25.00, 10, 3, "Prateleira A3"], start=1):
        ws.cell(row=2, column=i, value=v)

    ws.freeze_panes = "A2"

    # Validações nativas (bloqueiam erro antes do upload)
    dv_un = DataValidation(type="list", formula1='"%s"' % ",".join(UNIDADES), allow_blank=True,
                           errorTitle="Unidade inválida", error="Use: " + ", ".join(UNIDADES))
    ws.add_data_validation(dv_un)
    dv_un.add("C2:C%d" % (MAX_LINHAS + 1))
    for col in "DEFG":   # preços e estoques: número >= 0
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0,
                            errorTitle="Valor inválido", error="Use número igual ou maior que zero (ex: 12.50).")
        ws.add_data_validation(dv)
        dv.add("%s2:%s%d" % (col, col, MAX_LINHAS + 1))

    # Aba 2: como preencher
    ws2 = wb.create_sheet("Como preencher")
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 82
    guia = [
        ("Coluna", "Como preencher"),
        ("Codigo", "Obrigatório. 2 a 50 caracteres. É a CHAVE: código que já existe ATUALIZA a peça; novo CRIA."),
        ("Descricao", "Obrigatório. 2 a 200 caracteres."),
        ("Unidade", "UN, PC, CX, KG, LT, M ou ML. Em branco vira UN."),
        ("Preco Custo", "Número igual ou maior que zero. Ex: 12.50 (ou 12,50)."),
        ("Preco Venda", "Número igual ou maior que zero."),
        ("Estoque Atual", "Número igual ou maior que zero. Aceita 3 casas (ex: 1.5)."),
        ("Estoque Minimo", "Número igual ou maior que zero. Abaixo dele a peça aparece em alerta."),
        ("Localizacao", "Opcional. Onde a peça fica (ex: Prateleira A3)."),
        ("", ""),
        ("Linha de exemplo", "A linha com código EX-001 é só exemplo: ela é IGNORADA na importação. Pode apagar."),
        ("Repetição", "Subir o mesmo arquivo de novo não duplica: ele atualiza pelo código."),
    ]
    for r, (a, b) in enumerate(guia, start=1):
        ca, cb = ws2.cell(row=r, column=1, value=a), ws2.cell(row=r, column=2, value=b)
        if r == 1:
            for c in (ca, cb):
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor=VERDE_CADASTRO)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ──────────────────────────────────────────────────────────────────────────────
# Leitura e validação
# ──────────────────────────────────────────────────────────────────────────────

def _pt_br(v):
    """Célula numérica do Excel (float/int) -> texto pt-BR que o PecaIn espera.
    Texto fica como veio (o usuário pode ter digitado 1.234,56)."""
    if isinstance(v, bool):
        return ""
    if isinstance(v, (int, float)):
        return str(v).replace(".", ",")
    return "" if v is None else str(v).strip()


def _texto(v):
    return "" if v is None else str(v).strip()


def ler_linhas(arquivo) -> list[list]:
    """Lê o .xlsx e devolve as linhas de dados cruas (sem o cabeçalho).

    Levanta ValueError se o arquivo não for um Excel válido.
    """
    try:
        wb = load_workbook(arquivo, read_only=True, data_only=True)
    except Exception:
        raise ValueError("Arquivo não é um Excel (.xlsx) válido.")
    try:
        ws = wb["Peças"] if "Peças" in wb.sheetnames else wb.worksheets[0]
        linhas = []
        for linha in ws.iter_rows(min_row=2, max_col=len(COLUNAS), values_only=True):
            if all(c is None or str(c).strip() == "" for c in linha):
                continue   # linha em branco
            linhas.append(list(linha))
            if len(linhas) > MAX_LINHAS:
                raise ValueError(f"A planilha passou de {MAX_LINHAS} linhas. Divida em arquivos menores.")
        return linhas
    finally:
        wb.close()


def validar(linhas: list[list]) -> dict:
    """Valida as linhas cruas. Devolve {validas, problemas, resumo}.

    `validas`: [{codigo, descricao, unidade, acao('novo'|'atualiza'), linhas(dict pro banco)}]
    `problemas`: [{linha, codigo, erro}]
    Linha de exemplo (EX-001) é ignorada. Código repetido no arquivo vira problema.
    """
    validas, problemas = [], []
    vistos: dict[str, int] = {}
    for i, bruta in enumerate(linhas, start=2):   # linha 2 é a primeira de dados
        bruta = (bruta + [None] * len(COLUNAS))[: len(COLUNAS)]
        cod_cru = _texto(bruta[0])
        if cod_cru.upper() == COD_EXEMPLO:
            continue   # exemplo: ignora
        dados_form = {
            "codigo": cod_cru,
            "descricao": _texto(bruta[1]),
            "unidade": _texto(bruta[2]) or "UN",
            "preco_custo": _pt_br(bruta[3]),
            "preco_venda": _pt_br(bruta[4]),
            "estoque_atual": _pt_br(bruta[5]),
            "estoque_minimo": _pt_br(bruta[6]),
            "localizacao": _texto(bruta[7]) or None,
        }
        try:
            peca = PecaIn(**dados_form)
        except Exception as exc:
            msgs = exc.errors() if hasattr(exc, "errors") else [{"msg": str(exc)}]
            erro = "; ".join(str(m.get("msg", "")).replace("Value error, ", "") for m in msgs) or "Linha inválida."
            problemas.append({"linha": i, "codigo": cod_cru or "(sem código)", "erro": erro})
            continue
        if peca.codigo in vistos:
            problemas.append({"linha": i, "codigo": peca.codigo,
                              "erro": f"Código repetido na planilha (já na linha {vistos[peca.codigo]})."})
            continue
        vistos[peca.codigo] = i
        existe = query_one("SELECT id FROM peca WHERE codigo = ? AND D_E_L_E_T = 0", (peca.codigo,))
        validas.append({
            "codigo": peca.codigo, "descricao": peca.descricao, "unidade": peca.unidade,
            "acao": "atualiza" if existe else "novo", "linhas": peca.para_linhas(),
        })
    resumo = {
        "total": len(validas) + len(problemas),
        "novos": sum(1 for v in validas if v["acao"] == "novo"),
        "atualiza": sum(1 for v in validas if v["acao"] == "atualiza"),
        "erros": len(problemas),
    }
    return {"validas": validas, "problemas": problemas, "resumo": resumo}
